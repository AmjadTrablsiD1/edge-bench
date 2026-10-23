"""Core tests: readers, edge search, measurement, waveform, Excel.

Every rule is checked twice, by two methods that share no code:
  * hand-built cases with answers worked out on paper, and
  * random captures checked against the naive sample-by-sample search in synth.py.
Real Logic 2 exports (simulated Logic 8 / Pro 16, Logic 2.4.48) are checked
binary-against-CSV: two export formats of the same capture must agree.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from app import excel, measure, readers, waveform
from app.capture import Channel, Segment
from app.edges import count_in, find_edge
from app.readers import binary, csv_digital, sal
from tests import synth

FIX = Path(__file__).parent / "fixtures"


# ---------------------------------------------------------------- readers

def test_v0_roundtrip(tmp_path):
    t = [1e-6, 2.5e-6, 7e-6]
    synth.write_v0(tmp_path / "digital_3.bin", 1, 0.0, 1e-5, t)
    cap = binary.load(tmp_path / "digital_3.bin")
    ch = cap.channel(3)
    assert ch.segments[0].initial == 1
    assert list(ch.segments[0].times) == t
    # initial high -> first edge falls, then rises, then falls
    assert [ch.segments[0].is_rising(i) for i in range(3)] == [False, True, False]
    assert cap.format_version == 0 and cap.sample_rate is None


def test_v1_chunks_and_gap(tmp_path):
    synth.write_v1(tmp_path / "digital_0.bin",
                   [(0, 0.0, 1.0, [0.2, 0.4]), (1, 2.0, 3.0, [2.5])], rate=1e6)
    cap = binary.load(tmp_path)
    ch = cap.channel(0)
    assert len(ch.segments) == 2 and cap.sample_rate == 1e6 and cap.format_version == 1
    assert ch.level_at(0.3) == 1 and ch.level_at(1.5) is None and ch.level_at(2.6) == 0
    e = find_edge(ch, 0.5)                    # nothing left in chunk 1 -> chunk 2's edge
    assert e.time == 2.5 and not e.rising


def test_truncated_file_is_refused(tmp_path):
    p = tmp_path / "digital_0.bin"
    synth.write_v0(p, 0, 0, 1, [0.1, 0.2, 0.3])
    p.write_bytes(p.read_bytes()[:-4])
    with pytest.raises(binary.BinaryFormatError, match="incomplete"):
        binary.read_header(p)


def test_not_saleae_is_refused(tmp_path):
    p = tmp_path / "digital_0.bin"
    p.write_bytes(b"hello world, not a capture at all")
    with pytest.raises(binary.BinaryFormatError, match="SALEAE"):
        binary.read_header(p)


def test_unknown_file_type(tmp_path):
    p = tmp_path / "notes.txt"
    p.write_text("x")
    with pytest.raises(ValueError, match="not a Saleae capture"):
        readers.load(p)


@pytest.mark.parametrize("folder", ["logic8_bin", "pro16_bin"])
def test_real_logic2_binary_exports(folder):
    cap = binary.load(FIX / folder)
    n = 16 if "16" in folder else 8
    assert [c.index for c in cap.channels] == list(range(n))
    for c in cap.channels:
        s = c.segments[0]
        assert len(s) > 0 and np.all(np.diff(s.times) > 0)
        assert s.begin <= s.times[0] and s.times[-1] <= s.end


def test_binary_and_csv_of_the_same_capture_agree():
    """Two Logic 2 export formats of one capture: same edges, same directions."""
    b = binary.load(FIX / "logic8_bin")
    c = csv_digital.load(FIX / "logic8_digital.csv")
    assert [ch.index for ch in c.channels] == [ch.index for ch in b.channels]
    for cb, cc in zip(b.channels, c.channels):
        sb, sc = cb.segments[0], cc.segments[0]
        assert sb.initial == sc.initial
        # CSV prints 9 decimals -> 1 ns; Logic 8 at 25 MS/s samples every 40 ns
        assert np.allclose(sb.times, sc.times, atol=1e-9, rtol=0)


def test_sal_meta_without_logic2():
    m = sal.read_meta(FIX / "sim_logic8.sal")
    assert m.device == "Logic 8" and m.sample_rate == 25e6 and m.digital == list(range(8))
    assert m.names[0] == "Channel 0" and m.simulation
    m16 = sal.read_meta(FIX / "sim_pro16.sal")
    assert m16.digital == list(range(16)) and m16.sample_rate == 100e6


def test_sal_meta_rejects_non_zip(tmp_path):
    p = tmp_path / "broken.sal"
    p.write_bytes(b"PK\x03\x04 not really")
    with pytest.raises(ValueError, match="not a readable"):
        sal.read_meta(p)


# ------------------------------------------------------------- edge search

def _ch(initial, times, begin=0.0, end=1.0, idx=0, name="c"):
    return Channel(idx, name, [Segment(initial, begin, end, np.asarray(times, dtype=float))])


def test_find_edge_by_hand():
    ch = _ch(0, [0.1, 0.2, 0.3, 0.4])            # rise .1, fall .2, rise .3, fall .4
    assert find_edge(ch, 0.0).time == 0.1
    assert find_edge(ch, 0.15, "rising").time == 0.3
    assert find_edge(ch, 0.15, "falling").time == 0.2
    assert find_edge(ch, 0.0, "rising", skip=1).time == 0.3   # 2nd rising
    assert find_edge(ch, 0.0, "falling", skip=1).time == 0.4
    assert find_edge(ch, 0.0, "falling", skip=2) is None
    assert find_edge(ch, 0.2).time == 0.2                     # inclusive: same instant
    assert find_edge(ch, 0.2, inclusive=False).time == 0.3
    assert find_edge(ch, 0.41) is None


def test_glitch_filter_by_hand():
    us = 1e-6
    # Contact bounce on a rising edge: L -> H(10) L(10.2) H(10.5), then stable.
    ch = _ch(0, [10 * us, 10.2 * us, 10.5 * us, 50 * us], end=100 * us)
    e = find_edge(ch, 0.0, "rising", glitch_s=1 * us)
    assert e.time == pytest.approx(10.5 * us) and e.first_cross == pytest.approx(10 * us)
    assert e.bounce_s == pytest.approx(0.5 * us)
    # Without the filter the first crossing is the edge.
    assert find_edge(ch, 0.0, "rising").time == pytest.approx(10 * us)
    # A lone 100 ns spike on a low line is dropped; the real edge at 50 us is found.
    sp = _ch(0, [5 * us, 5.1 * us, 50 * us], end=100 * us)
    assert find_edge(sp, 0.0, "any", glitch_s=1 * us).time == pytest.approx(50 * us)
    assert find_edge(sp, 0.0, "any").time == pytest.approx(5 * us)


@pytest.mark.parametrize("seed", range(12))
@pytest.mark.parametrize("glitch", [0, 3])
def test_engine_matches_naive_sample_search(seed, glitch):
    """Random streams with glitches, every kind, many start times and occurrences."""
    rng = np.random.default_rng(seed)
    rate = 1e6
    s = synth.random_samples(rng, 4000, p_toggle=0.02, glitch_p=0.4, glitch_len=(1, 4))
    initial, times = synth.samples_to_edges(s, rate)
    ch = _ch(initial, times, begin=0.0, end=len(s) / rate)
    gl_s = (glitch - 0.5) / rate if glitch else 0.0
    for t_samp in rng.integers(0, len(s), 25):
        t = t_samp / rate
        for kind in ("any", "rising", "falling"):
            for skip in (0, 1, 3):
                want = synth.naive_first_edge(s, rate, t, kind, glitch, skip)
                got = find_edge(ch, t, kind, gl_s, skip)
                if want is None:
                    assert got is None, (t, kind, skip)
                else:
                    assert got is not None, (t, kind, skip, want)
                    assert got.time == pytest.approx(want[0], abs=1e-12)
                    assert got.rising == want[1]


# ------------------------------------------------------------- measurement

def _capture(tmp_path, spec, rate=1e6, end=1e-3):
    """spec: {index: (initial, [edge times])} -> a binary-export capture with known rate."""
    for idx, (initial, times) in spec.items():
        synth.write_v1(tmp_path / f"digital_{idx}.bin", [(initial, 0.0, end, times)], rate)
    return binary.load(tmp_path)


def test_measure_manual(tmp_path):
    us = 1e-6
    cap = _capture(tmp_path, {
        0: (0, [100 * us, 300 * us, 500 * us]),     # ref: rises at 100 and 500
        1: (1, [350 * us]),                           # falls 250 us after the 1st ref
        2: (0, []),                                   # never moves
        3: (0, [50 * us, 900 * us]),                  # rose before the ref; next edge at 900
    })
    s = measure.Settings(ref_channel=0, ref_edge="rising", target_edge="any")
    r = measure.run(cap, s)
    by = {c.index: c for c in r.channels}
    assert r.reference.time == pytest.approx(100 * us)
    assert by[0].status == "reference" and by[0].dt == 0
    assert by[1].status == "ok" and by[1].dt == pytest.approx(250 * us) and by[1].edge == "falling"
    assert by[2].status == "no_edge"
    assert by[3].status == "ok" and by[3].dt == pytest.approx(800 * us) and by[3].level_at_ref == 1
    assert r.resolution_s == pytest.approx(1 * us)

    # 2nd rising edge of the reference
    r2 = measure.run(cap, measure.Settings(ref_channel=0, occurrence=2))
    assert r2.reference.time == pytest.approx(500 * us)
    assert {c.index: c for c in r2.channels}[1].status == "no_edge"

    # A 300 us window turns channel 3's late edge into "out of window"
    r3 = measure.run(cap, measure.Settings(ref_channel=0, window_s=300 * us))
    c3 = {c.index: c for c in r3.channels}[3]
    assert c3.status == "out_of_window" and c3.dt is None and c3.next_dt == pytest.approx(800 * us)


def test_measure_auto_first_change(tmp_path):
    us = 1e-6
    cap = _capture(tmp_path, {
        0: (0, [400 * us]),
        1: (0, [120 * us]),          # first channel to change
        2: (1, [130 * us]),          # falls 10 us later
    })
    r = measure.run(cap, measure.Settings(ref_mode="auto", ref_edge="any"))
    assert r.reference.channel == 1 and r.reference.time == pytest.approx(120 * us)
    by = {c.index: c for c in r.channels}
    assert by[2].dt == pytest.approx(10 * us) and by[0].dt == pytest.approx(280 * us)
    # Excluding channel 1 makes channel 2 the first to change
    r2 = measure.run(cap, measure.Settings(ref_mode="auto", ref_edge="any", excluded=[1]))
    assert r2.reference.channel == 2
    assert {c.index: c for c in r2.channels}[1].status == "excluded"


def test_measure_simultaneous_edges_give_zero(tmp_path):
    cap = _capture(tmp_path, {0: (0, [2e-4]), 1: (0, [2e-4])})
    r = measure.run(cap, measure.Settings(ref_channel=0))
    assert {c.index: c for c in r.channels}[1].dt == 0.0


def test_measure_validates():
    with pytest.raises(ValueError):
        measure.Settings.from_dict({"ref_mode": "manual"})          # no channel
    with pytest.raises(ValueError):
        measure.Settings.from_dict({"ref_channel": 0, "ref_edge": "up"})
    with pytest.raises(ValueError):
        measure.Settings.from_dict({"ref_channel": 0, "occurrence": 0})


def test_measure_on_real_export_matches_csv():
    """Same capture, two formats, same answers."""
    s = measure.Settings(ref_channel=0, ref_edge="rising", occurrence=3)
    rb = measure.run(binary.load(FIX / "logic8_bin"), s)
    rc = measure.run(csv_digital.load(FIX / "logic8_digital.csv"), s)
    assert rb.reference.time == pytest.approx(rc.reference.time, abs=1e-9)
    for a, b in zip(rb.channels, rc.channels):
        assert a.status == b.status
        if a.dt is not None:
            assert a.dt == pytest.approx(b.dt, abs=2e-9)


# ---------------------------------------------------------------- waveform

def test_waveform_exact_and_buckets_agree_on_counts():
    rng = np.random.default_rng(1)
    times = np.sort(rng.uniform(0, 1, 20000))
    ch = _ch(0, times)
    cap = type("Cap", (), {"channels": [ch]})()
    few = waveform.decimate(cap, 0.5, 0.5001, 400)["channels"][0]["segments"][0]
    assert few["mode"] == "exact" and len(few["times"]) == count_in(ch, 0.5, 0.5001)
    many = waveform.decimate(cap, 0.0, 1.0, 400)["channels"][0]["segments"][0]
    assert many["mode"] == "buckets" and len(many["levels"]) == 400
    # a bucket with no edge keeps the level of the previous bucket's end
    lv, ct = many["levels"], many["counts"]
    for k in range(1, 400):
        if ct[k - 1] == 0:
            assert lv[k] == lv[k - 1]


# ------------------------------------------------------------------- Excel

def test_excel_roundtrip(tmp_path):
    us = 1e-6
    cap = _capture(tmp_path, {0: (0, [100 * us]), 1: (1, [350 * us]), 2: (0, [])})
    r = measure.run(cap, measure.Settings(ref_channel=0))
    out = excel.write([r, r], tmp_path / "out.xlsx", unit="µs")
    wb = load_workbook(out)
    ws = wb["Results"]
    hdr = [c.value for c in ws[1]]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    dt_col = hdr.index("Δt [µs]")
    assert rows[1][dt_col] == pytest.approx(250.0)
    assert rows[0][hdr.index("Status")] == "Reference"
    assert rows[2][hdr.index("Status")] == "No edge after reference" and rows[2][dt_col] is None
    assert "Batch matrix" in wb.sheetnames and "Settings" in wb.sheetnames
    m = [[c.value for c in row] for row in wb["Batch matrix"].iter_rows(min_row=2)]
    assert m[0][3] == pytest.approx(250.0)      # Capture, Reference, ch0, ch1


def test_excel_unit_scaling(tmp_path):
    cap = _capture(tmp_path, {0: (0, [1e-4]), 1: (0, [3.5e-4])})
    r = measure.run(cap, measure.Settings(ref_channel=0))
    for unit, want in (("ns", 250000.0), ("µs", 250.0), ("ms", 0.25), ("s", 2.5e-4)):
        ws = load_workbook(excel.write([r], tmp_path / f"{unit}.xlsx", unit=unit))["Results"]
        hdr = [c.value for c in ws[1]]
        assert ws.cell(row=3, column=hdr.index(f"Δt [{unit}]") + 1).value == pytest.approx(want)


def test_excel_rounds_to_resolution(tmp_path):
    """2.1884 ms measured at 10 ns resolution is 2188.4 µs, not 2188.3999999 µs."""
    t_ref, t_edge = 1.2e-3, 3.3884e-3
    cap = _capture(tmp_path, {0: (0, [t_ref]), 1: (0, [t_edge])}, rate=100e6, end=5e-3)
    r = measure.run(cap, measure.Settings(ref_channel=0))
    ws = load_workbook(excel.write([r], tmp_path / "r.xlsx", unit="µs"))["Results"]
    hdr = [c.value for c in ws[1]]
    assert ws.cell(row=3, column=hdr.index("Δt [µs]") + 1).value == 2188.4
    # and a value that is noisy in binary floating point
    assert excel._num(2188.3999999999996e-6, 1e-6, 10e-9) == 2188.4
    assert excel._num(2188.3999999999996e-6, 1e-6, None) == 2188.4      # 1 ns fallback

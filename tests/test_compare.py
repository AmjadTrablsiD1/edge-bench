"""Baseline vs new capture.

The power-up pair has differences planted on purpose (tests/fixtures/
make_power_up.py), so every verdict below is known before the code runs.  The
span metrics (duty, edge count, shortest pulse) are checked against a naive
sample-by-sample computation that shares no code with app/compare.py.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from app import excel
from app.capture import Channel, Segment
from app.compare import Tolerance, compare, span_activity
from app.measure import Settings
from app.readers import binary, csv_digital
from tests import synth

FIX = Path(__file__).parent / "fixtures"
BASE = FIX / "power_up_sequence.csv"
NEW = FIX / "power_up_new_fw.csv"


@pytest.fixture(scope="module")
def pair():
    return csv_digital.load(BASE), csv_digital.load(NEW)


def test_same_capture_is_same(pair):
    a, _ = pair
    c = compare(a, a, Settings(ref_channel=0))
    assert {r.verdict for r in c.rows} == {"same"}
    assert c.summary["same"] == 8 and c.summary["critical"] == 0


def test_planted_firmware_differences(pair):
    a, b = pair
    c = compare(a, b, Settings(ref_channel=0))
    v = {r.index: r.verdict for r in c.rows}
    assert v == {0: "same", 1: "changed", 2: "same", 3: "critical",
                 4: "same", 5: "changed", 6: "critical", 7: "changed"}
    rows = {r.index: r for r in c.rows}
    # RESET_N: +500 us, +25 %
    assert rows[1].dt_delta == pytest.approx(500e-6, abs=1e-9)
    assert [x.metric for x in rows[1].reasons] == ["dt"]       # not flagged twice via time high
    assert "Δt moved +500 µs (+25 %): 2 ms → 2.5 ms" in [x.text for x in rows[1].reasons]
    # PLL_LOCK never locks: edge gone and activity stopped
    kinds = {x.metric for x in rows[3].reasons if x.severity == "critical"}
    assert kinds == {"dt", "edges"}
    # SPI_SCK twice as fast: first pulse and shortest pulse halve, count and duty unchanged
    m5 = {x.metric for x in rows[5].reasons}
    assert m5 == {"first_pulse", "min_pulse"}
    assert rows[5].a.edges == rows[5].b.edges == 32
    # IRQ_N idles low: level at reference and direction both changed
    m6 = {x.metric for x in rows[6].reasons}
    assert {"level_at_ref", "edge"} <= m6
    # LED: shorter bounce spikes
    assert any("Shorter pulses" in x.text and "40 ns → 10 ns" in x.text for x in rows[7].reasons)


def test_swapping_a_and_b_keeps_the_verdicts(pair):
    a, b = pair
    ab = {r.index: r.verdict for r in compare(a, b, Settings(ref_channel=0)).rows}
    ba = {r.index: r.verdict for r in compare(b, a, Settings(ref_channel=0)).rows}
    # "shorter pulses" is one-directional by design (only shrinking pulses are risky)
    assert {k: v for k, v in ab.items() if k not in (5, 7)} == {k: v for k, v in ba.items() if k not in (5, 7)}


def _shifted(tmp_path, shift, rate=100e6):
    for name, extra in (("a", 0.0), ("b", shift)):
        d = tmp_path / name
        d.mkdir()
        synth.write_v1(d / "digital_0.bin", [(0, 0.0, 1e-2, [1e-3])], rate)
        synth.write_v1(d / "digital_1.bin", [(0, 0.0, 1e-2, [3e-3 + extra])], rate)
    return binary.load(tmp_path / "a"), binary.load(tmp_path / "b")


def test_tolerance_absolute_relative_and_resolution_floor(tmp_path):
    a, b = _shifted(tmp_path, 0.5e-6)                # 2 ms -> 2.0005 ms
    s = Settings(ref_channel=0)
    assert compare(a, b, s).rows[1].verdict == "same"                           # < 1 us and < 5 %
    tight = Tolerance(time_abs_s=0.1e-6, time_rel=0.0)
    assert compare(a, b, s, tight).rows[1].verdict == "changed"
    # both limits must be exceeded: 0.5 us is over 0.1 us but under 1 % of 2 ms
    assert compare(a, b, s, Tolerance(time_abs_s=0.1e-6, time_rel=0.01)).rows[1].verdict == "same"


def test_one_sample_jitter_is_never_a_change(tmp_path):
    a, b = _shifted(tmp_path, 10e-9)                 # one sample at 100 MS/s
    r = compare(a, b, Settings(ref_channel=0), Tolerance(time_abs_s=0, time_rel=0)).rows[1]
    assert r.verdict == "same" and r.dt_delta == pytest.approx(10e-9, abs=1e-12)


def test_missing_reference_is_reported(pair):
    a, b = pair
    c = compare(a, b, Settings(ref_channel=3))         # PLL_LOCK never rises in B
    assert c.reference_a is not None and c.reference_b is None
    assert any("No reference edge in capture B" in n for n in c.notes)


@pytest.mark.parametrize("seed", range(6))
def test_span_metrics_match_naive_samples(seed):
    """Duty = mean of the samples; edges and shortest pulse counted sample by sample."""
    rng = np.random.default_rng(seed)
    rate = 1e6
    s = synth.random_samples(rng, 5000, p_toggle=0.03, glitch_p=0.3)
    initial, times = synth.samples_to_edges(s, rate)
    ch = Channel(0, "x", [Segment(initial, 0.0, len(s) / rate, np.asarray(times))])
    k0, k1 = sorted(int(x) for x in rng.integers(1, len(s), 2))
    k1 = max(k1, k0 + 10)
    edges, duty, shortest = span_activity(ch, k0 / rate, k1 / rate)
    window = s[k0:k1]
    want_edges = [k for k in range(k0, k1 + 1) if k < len(s) and s[k] != s[k - 1]]
    assert edges == len(want_edges)
    assert duty == pytest.approx(window.mean(), abs=1e-9)
    if len(want_edges) >= 2:
        assert shortest == pytest.approx(min(np.diff(want_edges)) / rate, abs=1e-12)


def test_compare_workbook(pair, tmp_path):
    a, b = pair
    c = compare(a, b, Settings(ref_channel=0))
    out = excel.write_comparison(c, tmp_path / "cmp.xlsx", unit="µs")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Comparison", "Baseline (A)", "New (B)", "Settings"]
    ws = wb["Comparison"]
    hdr = [x.value for x in ws[1]]
    rows = {r[0].value: [x.value for x in r] for r in ws.iter_rows(min_row=2)}   # keyed by channel A
    reset = rows[1]
    assert reset[hdr.index("Verdict")] == "Changed"
    assert reset[hdr.index("Δt A [µs]")] == 2000 and reset[hdr.index("Δt B [µs]")] == 2500
    assert reset[hdr.index("B − A [µs]")] == 500
    assert rows[3][hdr.index("Verdict")] == "Critical"
    assert "Edge disappeared" in rows[3][hdr.index("Why")]


# ------------------------------------------------------------ match by name

REWIRED = FIX / "power_up_new_fw_rewired.csv"


def _csv(path: Path, chans: dict[str, tuple[int, list[float]]], end: float = 1e-2) -> Path:
    """A tiny Logic 2 digital CSV: {name: (initial level, edge times)}."""
    names = list(chans)
    ticks = sorted({round(t * 1e9) for _, ts in chans.values() for t in ts})
    state = [chans[n][0] for n in names]
    rows = ["Time [s]," + ",".join(names), "0.000000000," + ",".join(map(str, state))]
    for k in ticks:
        for i, n in enumerate(names):
            state[i] ^= [round(t * 1e9) for t in chans[n][1]].count(k) & 1
        rows.append(f"{k * 1e-9:.9f}," + ",".join(map(str, state)))
    rows.append(f"{end:.9f}," + ",".join(map(str, state)))
    path.write_text("\n".join(rows) + "\n")
    return path


def test_by_name_on_rewired_probes_gives_the_same_verdicts(pair):
    """Probes moved to other inputs: by name must equal the original by-number comparison."""
    a, b = pair
    rewired = csv_digital.load(REWIRED)
    want = {r.name_a: (r.verdict, r.dt_delta, [x.metric for x in r.reasons])
            for r in compare(a, b, Settings(ref_channel=0)).rows}
    got_cmp = compare(a, rewired, Settings(ref_channel=0), match="name")
    got = {r.name_a: (r.verdict, r.dt_delta, [x.metric for x in r.reasons]) for r in got_cmp.rows}
    assert got == want
    rows = {r.name_a: r for r in got_cmp.rows}
    assert (rows["RESET_N"].index, rows["RESET_N"].index_b, rows["RESET_N"].name_b) == (1, 3, "Reset_N")
    assert got_cmp.reference_b["name"] == "PWR_GOOD" and got_cmp.reference_b["channel"] == 1
    assert any("on a different channel in B" in n for n in got_cmp.notes)


def test_by_number_on_rewired_probes_is_wrong_and_says_so(pair):
    a, _ = pair
    c = compare(a, csv_digital.load(REWIRED), Settings(ref_channel=0))
    assert c.summary["same"] < 3                                    # garbage pairings
    assert any("try matching by name" in n for n in c.notes)


def test_exclusions_follow_the_name(pair):
    a, _ = pair
    c = compare(a, csv_digital.load(REWIRED), Settings(ref_channel=0, excluded=[3]), match="name")
    pll = next(r for r in c.rows if r.name_a == "PLL_LOCK")
    assert pll.verdict == "excluded" and pll.index_b == 7
    assert 7 in c.result_b.settings.excluded


def test_unmatched_and_duplicate_names(tmp_path):
    a = csv_digital.load(_csv(tmp_path / "a.csv", {"CLK": (0, [1e-3]), "DATA": (0, [2e-3]), "Data": (0, [3e-3])}))
    b = csv_digital.load(_csv(tmp_path / "b.csv", {"EN": (0, [5e-3]), "clk": (0, [1e-3]), "DATA": (0, [2e-3])}))
    c = compare(a, b, Settings(ref_channel=0), match="name")
    pairs = [(r.index, r.index_b, r.verdict) for r in c.rows]
    assert pairs == [(0, 1, "same"), (1, 2, "same"), (2, None, "missing"), (0, 0, "missing")]
    assert c.rows[2].reasons[0].text == "No channel with this name in B"
    assert c.rows[3].name_b == "EN" and c.rows[3].name_a is None
    assert any("paired in channel order" in n for n in c.notes)


def test_reference_name_missing_in_b(tmp_path):
    a = csv_digital.load(_csv(tmp_path / "a.csv", {"TRIG": (0, [1e-3]), "OUT": (0, [2e-3])}))
    b = csv_digital.load(_csv(tmp_path / "b.csv", {"START": (0, [1e-3]), "OUT": (0, [2e-3])}))
    c = compare(a, b, Settings(ref_channel=0), match="name")
    assert c.reference_a is not None and c.reference_b is None
    assert any("no channel named like A's reference (TRIG)" in n for n in c.notes)


def test_unknown_match_mode(pair):
    a, b = pair
    with pytest.raises(ValueError, match="channel matching"):
        compare(a, b, Settings(ref_channel=0), match="colour")


def test_workbook_records_name_matching(pair, tmp_path):
    a, _ = pair
    c = compare(a, csv_digital.load(REWIRED), Settings(ref_channel=0), match="name")
    wb = load_workbook(excel.write_comparison(c, tmp_path / "n.xlsx"))
    ws = wb["Comparison"]
    hdr = [x.value for x in ws[1]]
    reset = next([x.value for x in r] for r in ws.iter_rows(min_row=2) if r[2].value == "RESET_N")
    assert (reset[hdr.index("Channel A")], reset[hdr.index("Channel B")], reset[hdr.index("Name (B)")]) == (1, 3, "Reset_N")
    settings = {r[0].value: r[1].value for r in wb["Settings"].iter_rows(min_row=2)}
    assert settings["Matching"].startswith("by name")

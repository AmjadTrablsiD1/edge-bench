"""The Excel workbook: results only, never raw edges.

Sheets
  Results       one row per channel per capture, the reference row highlighted
  Batch matrix  (two or more captures) captures × channels, Δt in each cell
  Settings      every setting, the reference per capture, resolution, app version

Δt is written as a number in the chosen unit (µs by default), with the unit in
the header, so Excel can sort, filter and chart it straight away.
"""
from __future__ import annotations

import datetime as _dt
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.constants import C
from app.measure import Result

X = C.excel
_HEAD_FILL = PatternFill("solid", fgColor=X.header_fill)
_HEAD_FONT = Font(bold=True, color=X.header_font)
_REF_FILL = PatternFill("solid", fgColor=X.ref_fill)
_MISS_FILL = PatternFill("solid", fgColor=X.miss_fill)


def unit_scale(unit: str) -> float:
    for u in C.units.time:
        if u["label"] == unit:
            return u["scale"]
    raise ValueError(f"unknown time unit {unit!r}")


def _header(ws, cols: list[str], widths: list[int] | None = None) -> None:
    ws.append(cols)
    for i, _ in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = _HEAD_FILL, _HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = (widths or [])[i - 1] if widths else X.col_width
    ws.freeze_panes = "A2"


def _num(v: float | None, scale: float, resolution: float | None = None) -> float | None:
    """v in the chosen unit, rounded to what one sample period resolves
    (1 ns when the rate is unknown -- Logic 2's CSV precision).  Without this
    Excel shows binary-float noise such as 2188.3999999999 µs."""
    if v is None:
        return None
    res = resolution if resolution and resolution > 0 else 1e-9
    digits = max(0, min(12, math.ceil(math.log10(scale / res) - 1e-9)))
    return round(v / scale, digits)


def _results_sheet(ws, title: str, results: list[Result], unit: str) -> None:
    """One row per channel per capture; reference highlighted, misses marked."""
    sc = unit_scale(unit)
    labels = dict(C.measure.statuses)
    ws.title = title
    cols = ["Capture", "Channel", "Name", "Status", "Edge", f"Δt [{unit}]",
            "Edge time in capture [s]", f"Bounce before settling [{unit}]",
            f"Next edge (outside window) [{unit}]", "Level at reference", "Edges in capture"]
    w = [X.name_col_width, 9, X.name_col_width, 24, 10, X.col_width, 22, 18, 20, 12, 14]
    _header(ws, cols, w)
    for res in results:
        cap_name = Path(res.file).name
        for r in res.channels:
            rs = res.resolution_s
            ws.append([cap_name, r.index, r.name, labels.get(r.status, r.status), r.edge,
                       _num(r.dt, sc, rs), _num(r.time, 1.0, rs), _num(r.bounce, sc, rs),
                       _num(r.next_dt, sc, rs),
                       r.level_at_ref, r.edges_total])
            row = ws.max_row
            fill = _REF_FILL if r.status == "reference" else (
                _MISS_FILL if r.status in ("no_edge", "out_of_window", "no_data") else None)
            if fill:
                for c in range(1, len(cols) + 1):
                    ws.cell(row=row, column=c).fill = fill
            for c in (6, 8, 9):
                ws.cell(row=row, column=c).number_format = "0.0########"
            ws.cell(row=row, column=7).number_format = "0.000000000"
    ws.auto_filter.ref = ws.dimensions


def write(results: list[Result], out: Path, unit: str = C.units.default_export_unit) -> Path:
    if not results:
        raise ValueError("nothing to export: no measured captures")
    sc = unit_scale(unit)
    wb = Workbook()

    _results_sheet(wb.active, X.sheet_results, results, unit)

    # ---- Batch matrix ------------------------------------------------------
    if len(results) > 1:
        wm = wb.create_sheet(X.sheet_matrix)
        chans: dict[int, str] = {}
        for res in results:
            for r in res.channels:
                chans.setdefault(r.index, r.name)
        order = sorted(chans)
        _header(wm, ["Capture", "Reference"] + [f"{chans[i]} [{unit}]" for i in order],
                [X.name_col_width, X.name_col_width] + [X.col_width] * len(order))
        for res in results:
            by = {r.index: r for r in res.channels}
            ref = res.reference
            row = [Path(res.file).name, f"{ref.name} {ref.edge}" if ref else "none"]
            for i in order:
                r = by.get(i)
                row.append(_num(r.dt, sc, res.resolution_s) if r and r.status in ("ok", "reference") else None)
            wm.append(row)
            for c in range(3, len(row) + 1):
                cell = wm.cell(row=wm.max_row, column=c)
                cell.number_format = "0.0########"
                if cell.value is None:
                    cell.fill = _MISS_FILL

    # ---- Settings ----------------------------------------------------------
    st = wb.create_sheet(X.sheet_settings)
    _header(st, ["Setting", "Value"], [34, 60])
    s = results[0].settings
    rows = [
        ("Exported by", f"{C.app.name} {C.app.version}"),
        ("Exported at", _dt.datetime.now().isoformat(timespec="seconds")),
        ("Reference mode", "first channel that changes" if s.ref_mode == "auto" else "manual"),
        ("Reference channel", "auto" if s.ref_mode == "auto" else s.ref_channel),
        ("Reference edge", s.ref_edge),
        ("Reference occurrence", s.occurrence if s.ref_mode == "manual" else "first"),
        ("Search from [s]", "capture start" if s.search_from_s is None else s.search_from_s),
        ("Target edge", s.target_edge),
        ("Rule", "first edge at or after the reference (same sample counts, Δt = 0)"),
        (f"Window [{unit}]", "none" if s.window_s is None else _num(s.window_s, sc)),
        (f"Glitch filter [{unit}]", "off" if not s.glitch_s else _num(s.glitch_s, sc)),
        ("Excluded channels", ", ".join(map(str, s.excluded)) or "none"),
        ("Time unit", unit),
    ]
    for r in rows:
        st.append(list(r))
    st.append([])
    st.append(["Capture", "Reference · time in capture [s] · resolution ±[" + unit + "]"])
    for res in results:
        ref = res.reference
        st.append([Path(res.file).name,
                   (f"{ref.name} {ref.edge} #{ref.occurrence} at {ref.time:.9f} s" if ref else "no reference edge")
                   + (f" · ±{res.resolution_s / sc:.6g}" if res.resolution_s else " · resolution unknown")])
        for n in res.notes:
            st.append(["", n])

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ------------------------------------------------------------- comparison

_CHANGED_FILL = PatternFill("solid", fgColor=X.ref_fill)


def write_comparison(cmp, out: Path, unit: str = C.units.default_export_unit) -> Path:
    """Comparison sheet (verdict per channel, A, B, difference, why), then each
    capture's own results, then everything needed to reproduce it."""
    from app.units import fmt_pct, fmt_time

    sc = unit_scale(unit)
    res = cmp.resolution_s
    verdicts = dict(C.compare.verdicts)
    wb = Workbook()
    ws = wb.active
    ws.title = C.compare.sheet_compare
    cols = ["Channel", "Name (A)", "Name (B)", "Verdict",
            f"Δt A [{unit}]", f"Δt B [{unit}]", f"B − A [{unit}]", "Edge A", "Edge B",
            "Level@ref A", "Level@ref B", f"First pulse A [{unit}]", f"First pulse B [{unit}]",
            "Edges A", "Edges B", "Time high A", "Time high B",
            f"Shortest pulse A [{unit}]", f"Shortest pulse B [{unit}]", "Why"]
    widths = [9, 20, 20, 11] + [X.col_width] * 3 + [9, 9, 11, 11] + [X.col_width] * 2 + [9, 9, 11, 11] \
        + [X.col_width] * 2 + [80]
    _header(ws, cols, widths)
    for r in cmp.rows:
        a, b = r.a, r.b
        g = lambda m, k: getattr(m, k) if m is not None else None  # noqa: E731
        ws.append([
            r.index, r.name_a, r.name_b, verdicts.get(r.verdict, r.verdict),
            _num(g(a, "dt"), sc, res), _num(g(b, "dt"), sc, res), _num(r.dt_delta, sc, res),
            g(a, "edge"), g(b, "edge"), g(a, "level_at_ref"), g(b, "level_at_ref"),
            _num(g(a, "first_pulse"), sc, res), _num(g(b, "first_pulse"), sc, res),
            g(a, "edges"), g(b, "edges"),
            None if g(a, "duty") is None else round(a.duty, 4),
            None if g(b, "duty") is None else round(b.duty, 4),
            _num(g(a, "min_pulse"), sc, res), _num(g(b, "min_pulse"), sc, res),
            "; ".join(x.text for x in r.reasons) or None,
        ])
        row = ws.max_row
        fill = {"critical": _MISS_FILL, "changed": _CHANGED_FILL}.get(r.verdict)
        if fill:
            for c in range(1, 5):
                ws.cell(row=row, column=c).fill = fill
        for c in (5, 6, 7, 12, 13, 18, 19):
            ws.cell(row=row, column=c).number_format = "0.0########"
        for c in (16, 17):
            ws.cell(row=row, column=c).number_format = "0.0%"
        ws.cell(row=row, column=len(cols)).alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = ws.dimensions

    _results_sheet(wb.create_sheet(), C.compare.sheet_a, [cmp.result_a], unit)
    _results_sheet(wb.create_sheet(), C.compare.sheet_b, [cmp.result_b], unit)

    st = wb.create_sheet(X.sheet_settings)
    _header(st, ["Setting", "Value"], [34, 80])
    t, s = cmp.tolerance, cmp.settings

    def ref_text(ref):
        return (f"{ref['name']} {ref['edge']} #{ref['occurrence']} at {ref['time']:.9f} s in capture"
                if ref else "no reference edge")

    rows = [
        ("Exported by", f"{C.app.name} {C.app.version}"),
        ("Exported at", _dt.datetime.now().isoformat(timespec="seconds")),
        ("Baseline (A)", cmp.file_a), ("New (B)", cmp.file_b),
        ("Reference in A", ref_text(cmp.reference_a)), ("Reference in B", ref_text(cmp.reference_b)),
        ("Reference rule", f"{'first channel to change' if s.ref_mode == 'auto' else f'channel {s.ref_channel}'}, "
                           f"{s.ref_edge} edge #{s.occurrence}"),
        ("Target edge", s.target_edge),
        ("Glitch filter", fmt_time(s.glitch_s) if s.glitch_s else "off"),
        ("Span compared after each reference", fmt_time(cmp.span_s, res) if cmp.span_s else "—"),
        ("Resolution used", f"±{fmt_time(res)}" if res else "unknown"),
        ("Delay tolerance", f"±{fmt_time(t.time_abs_s)} or ±{fmt_pct(t.time_rel)}, whichever is larger"),
        ("Pulse-width tolerance", f"±{fmt_time(t.pulse_abs_s)} or ±{fmt_pct(t.pulse_rel)}, whichever is larger"),
        ("Edge-count tolerance", f"±{t.count_abs} edges or ±{fmt_pct(t.count_rel)}, whichever is larger"),
        ("Time-high tolerance", f"±{fmt_pct(t.duty_abs)} of the span"),
        ("Never flagged", f"differences up to {C.compare.resolution_floor_samples} sample periods"),
        ("Matching", "channels matched by number (D0 with D0)"),
    ]
    for r in rows:
        st.append(list(r))
    for n in cmp.notes:
        st.append(["Note", n])

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
